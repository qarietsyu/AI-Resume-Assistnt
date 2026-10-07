"""Resume ATS Score Checker - Streamlit + Gemini Flash."""
import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-2.5-flash"  # change here if you want another Flash model
MAX_RESUME_CHARS = 20000
MIN_RESUME_CHARS = 150


# ---------- Helpers ----------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from PDF, DOCX or TXT bytes."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str) -> str:
    jd_part = (
        f"JOB DESCRIPTION:\n{job_description.strip()}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: not provided. Evaluate for general ATS-friendliness."
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyze the resume below{" against the job description" if job_description.strip() else ""}.
Be honest and strict. Do not invent facts that are not in the resume.

Return ONLY valid JSON with exactly this structure:
{{
  "overall_score": <integer 0-100>,
  "section_scores": {{
    "formatting": <integer 0-100>,
    "keywords": <integer 0-100>,
    "experience": <integer 0-100>,
    "education": <integer 0-100>,
    "skills": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", "..."],
  "weaknesses": ["<string>", "..."],
  "missing_keywords": ["<string>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "issue": "<string>", "suggestion": "<string>"}}
  ],
  "rewritten_examples": [
    {{"original": "<line from resume>", "improved": "<stronger version>"}}
  ]
}}

{jd_part}
RESUME:
{resume_text}
"""


def parse_json_response(text: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences or extra text."""
    cleaned = (text or "").strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("The AI response was not valid JSON. Please try again.")


def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def normalize_result(raw: dict) -> dict:
    """Make sure every field exists and has the right type, so the UI never crashes."""
    section_keys = ["formatting", "keywords", "experience", "education", "skills", "readability"]
    sections = raw.get("section_scores") or {}

    def str_list(key):
        items = raw.get(key) or []
        return [str(i) for i in items if str(i).strip()] if isinstance(items, list) else []

    improvements = []
    for item in raw.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
    examples = []
    for item in raw.get("rewritten_examples") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            examples.append({"original": str(item["original"]), "improved": str(item["improved"])})

    return {
        "overall_score": clamp_score(raw.get("overall_score")),
        "section_scores": {k: clamp_score(sections.get(k)) for k in section_keys},
        "summary": str(raw.get("summary", "")),
        "strengths": str_list("strengths"),
        "weaknesses": str_list("weaknesses"),
        "missing_keywords": str_list("missing_keywords"),
        "improvements": improvements,
        "rewritten_examples": examples,
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text[:MAX_RESUME_CHARS], job_description[:8000]),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good - needs some work"
    if score >= 40:
        return "Fair - needs significant work"
    return "Poor - major improvements needed"


# ---------- UI ----------
def main():
    st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="centered")
    st.title("📄 Resume ATS Score Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at aistudio.google.com")
        else:
            st.success("API key loaded")
        st.markdown("Supported files: PDF, DOCX, TXT")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, gives a more accurate score)", height=150
    )

    if st.button("Analyze Resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(text) < MIN_RESUME_CHARS:
            st.error("Very little text was found. If your resume is a scanned image, "
                     "upload a text-based PDF or DOCX instead (ATS systems cannot read images either).")
            return
        try:
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(api_key, text, job_description)
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return

    result = st.session_state.get("result")
    if not result:
        return

    st.divider()
    score = result["overall_score"]
    st.metric("Overall ATS Score", f"{score} / 100", score_label(score))
    st.progress(score / 100)
    if result["summary"]:
        st.write(result["summary"])

    st.subheader("Section scores")
    cols = st.columns(3)
    for i, (name, value) in enumerate(result["section_scores"].items()):
        cols[i % 3].metric(name.title(), f"{value}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["None identified"]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for w in result["weaknesses"] or ["None identified"]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("🛠️ How to improve")
    order = {"High": 0, "Medium": 1, "Low": 2}
    for item in sorted(result["improvements"], key=lambda x: order.get(x["priority"], 3)):
        icon = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}.get(item["priority"], "⚪")
        with st.expander(f"{icon} {item['priority']}: {item['issue']}"):
            st.write(item["suggestion"])

    if result["rewritten_examples"]:
        st.subheader("✍️ Rewrite examples")
        for ex in result["rewritten_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.markdown("---")

    st.download_button("Download report (JSON)", json.dumps(result, indent=2),
                       file_name="ats_report.json", mime="application/json")
    st.caption("This score is an AI estimate, not the output of a real ATS.")


if __name__ == "__main__":
    main()
